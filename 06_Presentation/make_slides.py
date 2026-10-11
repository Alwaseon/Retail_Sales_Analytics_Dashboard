import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

DARK_BG = RGBColor(0x11, 0x18, 0x27)
CARD_BG = RGBColor(0x1F, 0x29, 0x37)
TEXT_WHITE = RGBColor(0xF9, 0xFA, 0xFB)
TEXT_MUTED = RGBColor(0x9C, 0xA3, 0xAF)
ACCENT_GREEN = RGBColor(0x10, 0xB9, 0x81)

def apply_theme(slide, title_text):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK_BG
    bg.line.fill.background()
    
    if title_text:
        txBox = slide.shapes.add_textbox(Inches(0.75), Inches(0.4), Inches(11.83), Inches(0.8))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text.upper()
        p.font.name = "Segoe UI"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GREEN

def add_bullet_points(slide, left, top, width, height, points):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = RGBColor(0x37, 0x41, 0x51)
    
    txBox = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    for i, pt_text in enumerate(points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "• " + pt_text
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(14)

# Slide 1: Title
slide1 = prs.slides.add_slide(blank_layout)
apply_theme(slide1, "")
txBox = slide1.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.33), Inches(3))
tf = txBox.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "SUPERSTORE RETAIL PERFORMANCE\nOPTIMIZATION DASHBOARD"
p.font.name = "Segoe UI"
p.font.size = Pt(40)
p.font.bold = True
p.font.color.rgb = TEXT_WHITE
p.alignment = PP_ALIGN.LEFT

p2 = tf.add_paragraph()
p2.text = "Final Project Defense | Emil Asqarov, Data Analyst"
p2.font.name = "Segoe UI"
p2.font.size = Pt(20)
p2.font.color.rgb = ACCENT_GREEN
p2.space_before = Pt(20)

# Slide 2: Core KPIs
slide2 = prs.slides.add_slide(blank_layout)
apply_theme(slide2, "Executive Core KPIs Summary")
kpis = [
    "Gross Revenue: $2,297,200.86",
    "Net Corporate Profit: $286,397.02",
    "Corporate Profit Margin: 12.47%",
    "Average Order Value (AOV): $458.61",
    "Active Customer Ledger: 793 Unique Clients",
    "Total Shipped Volume: 37,873 Units"
]
add_bullet_points(slide2, Inches(0.75), Inches(1.5), Inches(11.83), Inches(5.2), kpis)

# Slide 3: Weeks 1 & 2
slide3 = prs.slides.add_slide(blank_layout)
apply_theme(slide3, "Data Engineering & Validation (Weeks 1 & 2)")
w12 = [
    "Problem Statement: Raw transaction logs suffered from inconsistent regional locale configurations and critical date parsing errors.",
    "Data Cleaning Workflow: Utilized Power Query to resolve American/European date format syntax, handle duplicate indices, and cast uniform mathematical types.",
    "Useful Fields Added: Extracted calendar logic (order_year, order_month), and compiled automated margin metrics to stabilize core reports."
]
add_bullet_points(slide3, Inches(0.75), Inches(1.5), Inches(11.83), Inches(5.2), w12)

# Slide 4: Week 3
slide4 = prs.slides.add_slide(blank_layout)
apply_theme(slide4, "Relational Infrastructure & Database Load (Week 3)")
w3 = [
    "Database Architecture: Deployed a structured SQL schema mirroring clean enterprise tables to establish a verified data source.",
    "Data Engineering: Resolved localized numeric notation barriers (converted string commas '261,96' to numeric database dots '261.96') using PostgreSQL ALTER TABLE transformations.",
    "Query Execution: Loaded 9,994 transactional rows to enable sub-category and customer cohort segmentation via optimized SQL grouping scripts."
]
add_bullet_points(slide4, Inches(0.75), Inches(1.5), Inches(11.83), Inches(5.2), w3)

# Slide 5: Week 4
slide5 = prs.slides.add_slide(blank_layout)
apply_theme(slide5, "Advanced Discount Elasticity Insights (Week 4 Deep-Dive)")
w4 = [
    "DAX Implementation: Built a dedicated 'Discount Profit Margin' DAX measure to test profit health against aggressive pricing models.",
    "The Core Discovery: Proved that a 10%-20% discount structure maintains a healthy net return, but pushing past a 50% discount ceiling triggers exponential decay.",
    "Catastrophic Collapse: At 70% and 80% off, net margin drops to a devastating floor of nearly -150%, actively destroying corporate capital."
]
add_bullet_points(slide5, Inches(0.75), Inches(1.5), Inches(11.83), Inches(5.2), w4)

# Slide 6: Recommendations
slide6 = prs.slides.add_slide(blank_layout)
apply_theme(slide6, "Strategic Recommendations for Leadership")
recs = [
    "Discount Cap: Enforce a hard 30% ceiling on promotional discounts for large furniture and high-tech lines to avoid negative returns.",
    "Logistics Restructuring: Terminate blanket free bulk shipping policies for oversized sub-categories (Tables, Bookcases) to protect net margins.",
    "Loyalty Matrix: Implement custom account management systems for the top-tier High-Value corporate accounts that secure the baseline of net recurring revenue."
]
add_bullet_points(slide6, Inches(0.75), Inches(1.5), Inches(11.83), Inches(5.2), recs)

output_path = os.path.join("06_Presentation", "final_presentation.pptx")
prs.save(output_path)
print("SUCCESS: final_presentation.pptx has been successfully generated in 06_Presentation folder!")
